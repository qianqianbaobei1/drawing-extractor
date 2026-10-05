# -*- coding: utf-8 -*-
import unittest
import test_support  # noqa: F401

from extractor.schema import (
    Box,
    Circuit,
    Component,
    Evidence,
    EvidenceType,
    ExtractionResult,
    InferenceType,
    RawExtraction,
    Uncertainty,
    validate_field_evidence,
)
from extractor.checker import (
    BREAKER_MODEL_REGEX,
    CheckSeverity,
    check_result,
    check_result_issues,
    clean_rated_amp,
)
from extractor.catalog_reconciler import (
    DrawingCatalogReconciler,
    expand_panel_range,
)


class TestEvidenceAndCatalogReconciler(unittest.TestCase):

    def test_clean_rated_amp_no_crash(self):
        """验证额定电流鲁棒解析，杜绝 float('C63') 等非纯数字字符串崩溃。"""
        # 1. 常见工程带前缀/脱扣曲线字符串
        self.assertEqual(clean_rated_amp("C63"), 63.0)
        self.assertEqual(clean_rated_amp("D100"), 100.0)
        self.assertEqual(clean_rated_amp("B16"), 16.0)
        self.assertEqual(clean_rated_amp("C16/2P"), 16.0)
        self.assertEqual(clean_rated_amp("/C63/3P"), 63.0)
        self.assertEqual(clean_rated_amp("In=63A"), 63.0)
        self.assertEqual(clean_rated_amp("In: 100A"), 100.0)
        self.assertEqual(clean_rated_amp("63A"), 63.0)
        self.assertEqual(clean_rated_amp("100/3P"), 100.0)
        self.assertEqual(clean_rated_amp(63), 63.0)
        self.assertEqual(clean_rated_amp(32.5), 32.5)

        # 2. 异常输入不抛 ValueError，安全返回 None
        self.assertIsNone(clean_rated_amp("Unknown"))
        self.assertIsNone(clean_rated_amp(""))
        self.assertIsNone(clean_rated_amp(None))
        self.assertIsNone(clean_rated_amp("ABC/XYZ"))

    def test_breaker_model_regex_fix(self):
        """验证型号正则 [CD] 正确性，杜绝 [C|D] 匹配字面 '|'。"""
        # 合法型号
        self.assertTrue(bool(BREAKER_MODEL_REGEX.match("NM1-C63/3P")))
        self.assertTrue(bool(BREAKER_MODEL_REGEX.match("CDB6-D32/2P")))
        self.assertTrue(bool(BREAKER_MODEL_REGEX.match("iC65N-C16/1P".upper())))
        self.assertTrue(bool(BREAKER_MODEL_REGEX.match("NXB-63-C25/3P+N".upper())))

        # 包含字面 '|' 的畸变字符必须判定为非法
        self.assertFalse(bool(BREAKER_MODEL_REGEX.match("NM1-|63/3P")))
        self.assertFalse(bool(BREAKER_MODEL_REGEX.match("INVALID_BREAKER")))

    def test_expand_panel_range(self):
        """验证图纸目录配电箱范围展开器。"""
        # 批量展开 JX1~JX21 (根治 JX 漏柜关键)
        jx_list = expand_panel_range("JX1~JX21")
        self.assertEqual(len(jx_list), 21)
        self.assertEqual(jx_list[0], "JX1")
        self.assertEqual(jx_list[-1], "JX21")

        # 紧凑写法 1AL1~5
        al_list = expand_panel_range("1AL1~5")
        self.assertEqual(al_list, ["1AL1", "1AL2", "1AL3", "1AL4", "1AL5"])

        # 短横线写法 2AP1-4
        ap_list = expand_panel_range("2AP1-4")
        self.assertEqual(ap_list, ["2AP1", "2AP2", "2AP3", "2AP4"])

        # 离散列表 1AL1, 1AL2, 2AP1
        multi = expand_panel_range("1AL1, 1AL2, 2AP1")
        self.assertEqual(multi, ["1AL1", "1AL2", "2AP1"])

    def test_catalog_reconciliation_missing_jx(self):
        """模拟 JX1~JX21 漏柜场景：图纸目录声明存在，但提取结果完全未进流水线。"""
        catalog_lines = [
            "01B-01 配电箱系统图(一) 1AL1~1AL3",
            "01B-02 动力配电箱系统图(二) 1AP1~1AP2",
            "01B-03 动力配电箱系统图(三) JX1~JX21",
        ]
        catalog_items = [
            DrawingCatalogReconciler.parse_catalog_line(line)
            for line in catalog_lines
        ]
        catalog_items = [item for item in catalog_items if item is not None]

        # 假设实际只提取到了 1AL1~1AL3 和 1AP1~1AP2，JX1~21 彻底未提取
        extracted_boxes = [
            Box(code="1AL1", name="照明箱1"),
            Box(code="1AL2", name="照明箱2"),
            Box(code="1AL3", name="照明箱3"),
            Box(code="1AP1", name="动力箱1"),
            Box(code="1AP2", name="动力箱2"),
        ]

        recon = DrawingCatalogReconciler.reconcile(catalog_items, extracted_boxes)
        self.assertTrue(recon.has_catalog)
        self.assertEqual(recon.total_declared_panels, 26)  # 3 + 2 + 21
        self.assertEqual(recon.covered_count, 5)
        self.assertEqual(recon.missing_count, 21)
        self.assertAlmostEqual(recon.coverage_rate, 5 / 26, places=3)
        self.assertEqual(len(recon.missing_box_codes), 21)

        # 验证生成的 ERROR 级别告警
        issues = DrawingCatalogReconciler.generate_reconciliation_issues(recon)
        self.assertTrue(len(issues) >= 1)
        jx_issue = [i for i in issues if "01B-03" in i.detail][0]
        self.assertEqual(jx_issue.severity, "ERROR")
        self.assertIn("【范围缺失】", jx_issue.detail)
        self.assertIn("JX1", jx_issue.detail)

    def test_field_evidence_policy(self):
        """验证字段级证据政策：严禁关键字段纯 MODEL_INFERENCE 编造。"""
        # 1. 允许物理文字
        valid_loc, _ = validate_field_evidence("box.location", EvidenceType.TEXT.value)
        self.assertTrue(valid_loc)

        # 2. 禁止安装位置臆测（防编造锅炉房/制丝工房）
        invalid_loc, msg_loc = validate_field_evidence("box.location", InferenceType.MODEL.value)
        self.assertFalse(invalid_loc)
        self.assertIn("违反证据政策", msg_loc)

        # 3. 禁止断路器无据推测
        invalid_brk, msg_brk = validate_field_evidence("circuit.breaker", InferenceType.MODEL.value)
        self.assertFalse(invalid_brk)

        # 4. 允许负荷用途语义推测
        valid_load, _ = validate_field_evidence("circuit.load_name", InferenceType.MODEL.value)
        self.assertTrue(valid_load)

    def test_check_result_issues_with_evidence_and_catalog(self):
        """测试 check_result_issues 结构化门禁及三级分类输出。"""
        res = ExtractionResult(
            boxes=[Box(code="1AL1", quantity=1)],
            circuits=[
                Circuit(box="1AL1", circuit_no="WL1", breaker="C16/1P", phase="L1"),
                Circuit(box="NON_EXIST_BOX", circuit_no="WL2", breaker="C20/1P", phase="L1"),
            ],
            components=[Component(name="小型断路器", spec="C16/1P", quantity=1, unit="只")],
            evidence_store={
                "box.location.1": Evidence(
                    evidence_id="box.location.1",
                    evidence_type=InferenceType.MODEL.value,
                    raw_content="制丝工房",
                )
            }
        )

        issues = check_result_issues(res)
        severities = {issue.severity for issue in issues}
        self.assertIn(CheckSeverity.ERROR.value, severities)

        # 检查是否精准拦截了证据违规与箱体不存在
        error_rules = {issue.rule_code for issue in issues if issue.severity == CheckSeverity.ERROR.value}
        self.assertIn("BOX_NOT_FOUND", error_rules)
        self.assertIn("EVIDENCE_POLICY_VIOLATION", error_rules)

    def test_fail_closed_grounded_field(self):
        """验证 GroundedField 默认值必须是 Fail-Closed (confidence is None, status is UNASSESSED)。"""
        from extractor.schema import GroundedField, ReviewStatus
        gf = GroundedField()
        self.assertIsNone(gf.confidence)
        self.assertEqual(gf.review_status, ReviewStatus.UNASSESSED.value)

    def test_structured_normalizer(self):
        """验证 Normalizer 能够正确参数化复杂线缆与断路器字符串。"""
        from extractor.normalizer import parse_cable, parse_breaker

        # 1. 复杂复合电缆
        c1 = parse_cable("WDZ-BYJ-3x2.5+E2.5 -SC20-CC")
        self.assertEqual(c1.family, "WDZ-BYJ")
        self.assertEqual(c1.core_count, 3)
        self.assertEqual(c1.section_mm2, 2.5)
        self.assertEqual(c1.pe_section_mm2, 2.5)
        self.assertIn("SC20", c1.laying_method)

        # 2. 动力大截面电缆
        c2 = parse_cable("ZR-YJV-4x35+1x16 CT")
        self.assertEqual(c2.family, "ZR-YJV")
        self.assertEqual(c2.core_count, 4)
        self.assertEqual(c2.section_mm2, 35.0)
        self.assertEqual(c2.pe_section_mm2, 16.0)

        # 3. 常见微断
        b1 = parse_breaker("iC65N-C16/2P 6kA")
        self.assertEqual(b1.manufacturer, "施耐德")
        self.assertEqual(b1.series, "iC65N")
        self.assertEqual(b1.curve, "C")
        self.assertEqual(b1.rated_current, 16.0)
        self.assertEqual(b1.poles, "2P")
        self.assertEqual(b1.breaking_capacity, "6kA")

        # 4. 塑壳带漏电
        b2 = parse_breaker("NM1-125S/3300 100A 3P 30mA")
        self.assertEqual(b2.manufacturer, "正泰")
        self.assertEqual(b2.rated_current, 100.0)
        self.assertEqual(b2.poles, "3P")
        self.assertEqual(b2.leakage_ma, 30)


    def test_assemble_populates_structured_parameters_and_claims(self):
        """验证 assemble 组装时自动清洗提取结构化参数，并建立字段级 Claim 字典。"""
        from extractor.assemble import assemble
        from extractor.schema import RawExtraction, ReviewStatus

        raw = RawExtraction(
            boxes=[Box(code="1AL1", name="照明配电箱", location="地下车库")],
            circuits=[
                Circuit(
                    box="1AL1",
                    circuit_no="WL1",
                    breaker="iC65N-C16/2P 6kA",
                    cable="WDZ-BYJ-3x2.5+E2.5 -SC20-CC",
                    phase="L1",
                    load_name="车道照明",
                )
            ],
            extra_devices=[],
            requirements=[],
            uncertainties=[],
        )
        res = assemble(raw)
        self.assertEqual(len(res.boxes), 1)
        self.assertEqual(len(res.circuits), 1)

        # 验证 Box Claims (未有人工确认或物理锚定时，默认为 PARSED_OK 或 UNASSESSED，严禁自封 CONFIRMED)
        box = res.boxes[0]
        self.assertIn("box.code", box.claims)
        self.assertEqual(box.claims["box.code"].review_status, ReviewStatus.PARSED_OK.value)
        self.assertIn("box.location", box.claims)
        self.assertEqual(box.claims["box.location"].review_status, ReviewStatus.UNASSESSED.value)

        # 验证 Circuit 结构化清洗参数
        c = res.circuits[0]
        self.assertIsNotNone(c.structured_breaker)
        self.assertEqual(c.structured_breaker.get("manufacturer"), "施耐德")
        self.assertEqual(c.structured_breaker.get("series"), "iC65N")
        self.assertEqual(c.structured_breaker.get("rated_current"), 16.0)

        self.assertIsNotNone(c.structured_cable)
        self.assertEqual(c.structured_cable.get("family"), "WDZ-BYJ")
        self.assertEqual(c.structured_cable.get("core_count"), 3)
        self.assertEqual(c.structured_cable.get("section_mm2"), 2.5)

        # 验证 Circuit Claims
        self.assertIn("circuit.circuit_no", c.claims)
        self.assertEqual(c.claims["circuit.circuit_no"].review_status, ReviewStatus.PARSED_OK.value)
        self.assertIn("circuit.breaker", c.claims)
        self.assertEqual(c.claims["circuit.breaker"].review_status, ReviewStatus.PARSED_OK.value)

    def test_sync_result_issues_preserves_three_tier_severities(self):
        """验证 _sync_result_issues 将 CheckIssue 精准映射至 uncertainties 并保留 ERROR/WARNING/INFO。"""
        from app import _sync_result_issues

        res = ExtractionResult(
            boxes=[Box(code="1AL1", quantity=1)],
            circuits=[
                Circuit(box="NON_EXIST_BOX", circuit_no="WL1", breaker="C16/1P"),
            ],
            components=[Component(name="小型断路器", spec="C16/1P", quantity=1, unit="只")],
        )
        _sync_result_issues(res)
        severities = {u.severity for u in res.uncertainties}
        self.assertIn("ERROR", severities)
        error_items = [u for u in res.uncertainties if u.severity == "ERROR"]
        self.assertTrue(any("箱体清单中未找到" in u.detail for u in error_items))

    def test_build_workbook_with_reconciliation_sheet(self):
        """验证 build_workbook 导出时自动写入“图纸目录对账审计”工作表。"""
        import tempfile
        import openpyxl
        from extractor.excel import build_workbook
        from extractor.catalog_reconciler import DrawingCatalogReconciler

        catalog_data = [
            {"sheet_no": "01B-01", "sheet_title": "1AL1配电箱系统图"},
            {"sheet_no": "01B-02", "sheet_title": "1AL2配电箱系统图"},
        ]
        reconciler = DrawingCatalogReconciler.from_records(catalog_data)
        recon_result = reconciler.reconcile(["1AL1"])

        res = ExtractionResult(
            title="测试工程报价",
            boxes=[Box(code="1AL1", quantity=1)],
            circuits=[Circuit(box="1AL1", circuit_no="WL1", breaker="C16/1P")],
            components=[Component(name="微型断路器", spec="C16/1P", quantity=1, unit="只")],
            reconciliation=recon_result,
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            temp_path = tf.name

        try:
            build_workbook(res, "测试副标题", temp_path, layout="3_sheets")
            wb = openpyxl.load_workbook(temp_path)
            self.assertIn("图纸目录对账审计", wb.sheetnames)
            ws = wb["图纸目录对账审计"]
            # 校验单元格内容
            self.assertIn("图纸目录对账审计表", ws.cell(row=1, column=1).value)
            self.assertEqual(ws.cell(row=4, column=2).value, "01B-01")
            self.assertEqual(ws.cell(row=5, column=2).value, "01B-02")
        finally:
            import os
            if os.path.exists(temp_path):
                os.remove(temp_path)

    def test_cad_texts_catalog_reconciliation_flow(self):
        """测试 CAD 文字解析 -> 目录对账 -> 异常识别全流程。"""
        from extractor.catalog_reconciler import DrawingCatalogReconciler
        from extractor.checker import check_result_issues

        cad_texts = [
            {"text": "图纸目录", "page": 1},
            {"text": "01B-01 1AL1 照明配电箱系统图", "page": 1},
            {"text": "01B-02 动力配电箱系统图(一) JX1~JX3", "page": 1},
            {"text": "01B-03 2AP1~2 空调配电箱系统图", "page": 1},
        ]
        parsed_items = []
        for t in cad_texts:
            item = DrawingCatalogReconciler.parse_catalog_line(t["text"])
            if item and item.declared_panels:
                parsed_items.append(item)

        self.assertEqual(len(parsed_items), 3)
        # 实际提取到的箱体仅有 1AL1 和 JX1
        extracted_boxes = [Box(code="1AL1", quantity=1), Box(code="JX1", quantity=1)]

        recon = DrawingCatalogReconciler.reconcile(parsed_items, extracted_boxes, source_name="CAD图纸目录")
        self.assertTrue(recon.has_catalog)
        self.assertEqual(recon.total_declared_panels, 6)  # 1AL1, JX1, JX2, JX3, 2AP1, 2AP2
        self.assertEqual(recon.covered_count, 2)
        self.assertEqual(recon.missing_count, 4)
        self.assertIn("JX2", recon.missing_box_codes)
        self.assertIn("2AP1", recon.missing_box_codes)

        # 验证门禁拦截器
        result = ExtractionResult(
            title="测试工程",
            boxes=extracted_boxes,
            circuits=[Circuit(box="1AL1", circuit_no="WL1", breaker="C16/1P")],
            components=[Component(name="微型断路器", spec="C16/1P", quantity=1, unit="只")],
            reconciliation=recon,
        )
        issues = check_result_issues(result)
        # 必须存在 ERROR 阻断项
        error_issues = [i for i in issues if i.severity == "ERROR"]
        self.assertTrue(any("整张图幅未纳入" in i.detail or "疑似系统图图幅漏传" in i.detail for i in error_issues))

    def test_raw_extraction_with_catalog_assembles_cleanly(self):
        """测试包含 catalog_items 的 RawExtraction 在 assemble 时自动贯通对账。"""
        from extractor.assemble import assemble
        from extractor.catalog_reconciler import DrawingCatalogReconciler

        catalog_items = [
            DrawingCatalogReconciler.parse_catalog_line("01B-01 1AL1配电箱系统图"),
            DrawingCatalogReconciler.parse_catalog_line("01B-02 1AL2配电箱系统图"),
        ]
        catalog_items = [ci for ci in catalog_items if ci]

        raw = RawExtraction(
            boxes=[Box(code="1AL1", name="照明箱", quantity=1)],
            circuits=[Circuit(box="1AL1", circuit_no="WL1", breaker="iC65N-C16/1P")],
            extra_devices=[],
            requirements=[],
            uncertainties=[],
            catalog_items=catalog_items,
        )
        res = assemble(raw)
        self.assertIsNotNone(res.reconciliation)
        self.assertTrue(res.reconciliation.has_catalog)
        self.assertEqual(res.reconciliation.covered_count, 1)
        self.assertEqual(res.reconciliation.missing_count, 1)
        self.assertEqual(res.reconciliation.missing_box_codes, ["1AL2"])

    def test_breaker_amp_suffix_cleaning(self):
        """测试脱扣特性后带 A 单位（如 C16A/1P、C10A）时额定电流与曲线的提纯及系列名过滤。"""
        from extractor.normalizer import parse_breaker
        from extractor.checker import clean_rated_amp, BREAKER_MODEL_REGEX

        # 1. 验证 normalizer.parse_breaker
        b1 = parse_breaker("C16A/1P")
        self.assertEqual(b1.curve, "C")
        self.assertEqual(b1.rated_current, 16.0)
        self.assertEqual(b1.poles, "1P")
        self.assertEqual(b1.series, "")

        b2 = parse_breaker("C10A")
        self.assertEqual(b2.curve, "C")
        self.assertEqual(b2.rated_current, 10.0)
        self.assertEqual(b2.series, "")

        b3 = parse_breaker("DZ47-63 C10A 1P")
        self.assertEqual(b3.curve, "C")
        self.assertEqual(b3.rated_current, 10.0)
        self.assertEqual(b3.poles, "1P")
        self.assertEqual(b3.series, "DZ47-63")

        b4 = parse_breaker("NXB-63-D32A/3P")
        self.assertEqual(b4.curve, "D")
        self.assertEqual(b4.rated_current, 32.0)
        self.assertEqual(b4.poles, "3P")
        self.assertEqual(b4.series, "NXB-63")

        # 2. 验证 checker.clean_rated_amp 与 BREAKER_MODEL_REGEX
        self.assertEqual(clean_rated_amp("C16A/1P"), 16.0)
        self.assertEqual(clean_rated_amp("MCB-63/C16A/1P"), 16.0)
        self.assertEqual(clean_rated_amp("C10A"), 10.0)
        self.assertEqual(clean_rated_amp("DZ47-63 C10A 1P"), 10.0)
        self.assertEqual(clean_rated_amp("NXB-63-D32A/3P"), 32.0)
        self.assertTrue(bool(BREAKER_MODEL_REGEX.match("MCB-63-C16A/1P")))

    def test_multi_box_cascade_coordination_isolation(self):
        """测试多箱体进出线开关级配按箱体独立核验，杜绝全局混淆虚假越级警告。"""
        from extractor.checker import check_result_issues

        # 箱体 1AP1: 进线 400A，出线 WP1 160A (不越级)
        # 箱体 1AL1: 进线 63A，出线 WL1 16A (不越级)
        # 全局混淆时，WP1 160A 会误与 1AL1 的 63A 比较而虚假越级报警
        res = ExtractionResult(
            boxes=[Box(code="1AP1"), Box(code="1AL1")],
            circuits=[
                Circuit(box="1AP1", circuit_no="进线", breaker="NM1-400/3P 400A", load_name="总进线"),
                Circuit(box="1AP1", circuit_no="WP1", breaker="NM1-160/3P 160A", load_name="照明分箱供电"),
                Circuit(box="1AL1", circuit_no="进线", breaker="MCB-63/C63/3P", load_name="箱进线"),
                Circuit(box="1AL1", circuit_no="WL1", breaker="C16/1P", load_name="照明1"),
            ],
            components=[
                Component(name="断路器（其他）", spec="NM1-400/3P 400A", quantity=1),
                Component(name="断路器（其他）", spec="NM1-160/3P 160A", quantity=1),
                Component(name="微型断路器", spec="MCB-63/C63/3P", quantity=1),
                Component(name="微型断路器", spec="C16/1P", quantity=1),
            ],
        )

        issues = check_result_issues(res)
        cascade_issues = [i for i in issues if i.rule_code == "CASCADE_OVERCURRENT"]
        self.assertEqual(len(cascade_issues), 0, "箱体级配独立核验时不应发生跨箱体虚假越级误报")

        # 真实越级测试：向 1AL1 增加 100A 支路，应精准指出 1AL1 越级
        res.circuits.append(Circuit(box="1AL1", circuit_no="WL2", breaker="100A/3P", load_name="超负荷设备"))
        res.components.append(Component(name="微型断路器", spec="100A/3P", quantity=1))
        issues_with_err = check_result_issues(res)
        cascade_errs = [i for i in issues_with_err if i.rule_code == "CASCADE_OVERCURRENT"]
        self.assertEqual(len(cascade_errs), 1)
        self.assertIn("1AL1", cascade_errs[0].detail)
        self.assertIn("WL2", cascade_errs[0].detail)

    def test_job_excel_export_retains_reconciliation(self):
        """测试从 job['data'] 字典重建 ExtractionResult 并导出 Excel 时完整保留图纸目录对账审计工作表。"""
        import tempfile
        import openpyxl
        from extractor.excel import build_workbook

        job_data = {
            "title": "对账测试工程",
            "boxes": [{"code": "1AL1", "name": "照明箱", "quantity": 1}],
            "circuits": [{"box": "1AL1", "circuit_no": "WL1", "breaker": "C16/1P"}],
            "components": [{"name": "微型断路器", "spec": "C16/1P", "quantity": 1, "unit": "只"}],
            "requirements": [],
            "uncertainties": [],
            "topology": [],
            "reconciliation": {
                "has_catalog": True,
                "catalog_source": "CAD图纸目录",
                "total_declared_panels": 2,
                "covered_count": 1,
                "missing_count": 1,
                "coverage_rate": 0.5,
                "missing_box_codes": ["1AL2"],
                "items": [
                    {"sheet_no": "01", "sheet_title": "1AL1系统图", "declared_panels": ["1AL1"], "matched_panels": ["1AL1"], "missing_panels": [], "status": "COVERED"},
                    {"sheet_no": "02", "sheet_title": "1AL2系统图", "declared_panels": ["1AL2"], "matched_panels": [], "missing_panels": ["1AL2"], "status": "MISSING"}
                ]
            }
        }

        # 模拟后端 /api/jobs/{job_id}/excel 中的重构与生成行为
        recon_data = job_data.get("reconciliation")
        topo_data = job_data.get("topology", [])
        result = ExtractionResult(
            title=job_data.get("title"),
            boxes=job_data.get("boxes", []),
            circuits=job_data.get("circuits", []),
            components=job_data.get("components", []),
            requirements=job_data.get("requirements", []),
            uncertainties=job_data.get("uncertainties", []),
            topology=topo_data,
            reconciliation=recon_data,
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            temp_path = tf.name

        try:
            build_workbook(result, "对账测试导出", temp_path, layout="3_sheets")
            wb = openpyxl.load_workbook(temp_path)
            self.assertIn("图纸目录对账审计", wb.sheetnames)
            ws = wb["图纸目录对账审计"]
            self.assertIn("目录声明总数: 2 台", ws.cell(row=2, column=1).value)
        finally:
            import os
            if os.path.exists(temp_path):
                os.remove(temp_path)

    def test_post_catalog_api_reconcile(self):
        """验证 POST /api/jobs/{job_id}/catalog 接口对账与持久化。"""
        from fastapi.testclient import TestClient
        from app import app
        import db

        client = TestClient(app)
        job_id = "test_catalog_api_job"
        job_payload = {
            "job_id": job_id,
            "id": job_id,
            "status": "done",
            "tenant_id": "default",
            "data": {
                "boxes": [{"code": "1AP1", "name": "动力箱1"}, {"code": "1AP2", "name": "动力箱2"}],
                "circuits": [],
                "components": [],
                "requirements": [],
                "uncertainties": [],
            }
        }
        db.db_save_job(job_payload)
        import app as app_mod
        app_mod.jobs[job_id] = job_payload

        try:
            # 提交 CSV / 文本格式目录，声明 1AP1~1AP4 (共4台，缺2台)
            payload = {
                "catalog_text": "01B-01 动力系统图(一) 1AP1~1AP4\n01B-02 照明系统图(二) 1AL1",
                "source_name": "人工补录目录清单"
            }
            res = client.post(f"/api/jobs/{job_id}/catalog", json=payload)
            self.assertEqual(res.status_code, 200)
            data = res.json()
            recon = data.get("data", {}).get("reconciliation")
            self.assertIsNotNone(recon)
            self.assertEqual(recon["total_declared_panels"], 5)
            self.assertEqual(recon["covered_count"], 2)
            self.assertEqual(recon["missing_count"], 3)

            # 通过 GET 端点核实
            get_res = client.get(f"/api/jobs/{job_id}/catalog")
            self.assertEqual(get_res.status_code, 200)
            get_recon = get_res.json()
            self.assertEqual(get_recon["missing_count"], 3)
            self.assertEqual(get_recon["catalog_source"], "人工补录目录清单")
        finally:
            db.db_clean_all_test_data()

    def test_excel_no_fabrication_location_and_ggd(self):
        """验证 Excel 导出禁止捏造位置 '配电间/动力间' 与猜柜型 'GGD(落地)'。"""
        import openpyxl
        import tempfile
        import os
        from extractor.excel import build_workbook

        box = Box(code="AP1", name="总箱", install="落地安装", location="", size="")
        cir = Circuit(box="AP1", circuit_no="N1", breaker="C16/1P", cable="BV-3x2.5", load_name="插座")
        result = ExtractionResult(
            title="防编造测试",
            boxes=[box],
            circuits=[cir],
            components=[],
            requirements=[],
            uncertainties=[],
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            path = tf.name

        try:
            build_workbook(result, "subtitle", path, layout="3_sheets")
            wb = openpyxl.load_workbook(path)
            # 1. 检查屏柜汇总表：无尺寸时绝不猜测为 GGD(落地)
            ws_sum = wb["屏柜汇总表"]
            box_model_val = ws_sum.cell(row=8, column=4).value  # D列 型号规格
            self.assertNotEqual(box_model_val, "GGD(落地)")
            self.assertFalse(box_model_val)  # 应该为空字符串

            # 2. 检查屏柜汇总表：无位置时备注写入 '待确认'，绝不能臆造为 '配电间/动力间'
            loc_val = ws_sum.cell(row=8, column=9).value  # I列 使用部位/所属楼栋车间
            self.assertEqual(loc_val, "待确认")
            self.assertNotIn("配电间/动力间", str(loc_val))
        finally:
            if os.path.exists(path):
                os.remove(path)

    def test_excel_brand_and_floor_prices_unlocked(self):
        """验证 Excel 导出的品牌贯穿、下限价解禁、DZ47非SPD及备注记录。"""
        import openpyxl
        import tempfile
        import os
        from extractor.excel import build_workbook

        box = Box(code="AL1", name="照明箱", install="挂墙", location="车间A", size="400*500*160")
        cir_in = Circuit(box="AL1", circuit_no="进线", breaker="SW-32A/2P", cable="", load_name="市电进线")
        # 带有通用微断 DZ47-63
        cir_out = Circuit(box="AL1", circuit_no="WL1", breaker="DZ47-63 C10/1P", cable="BV-2x2.5", load_name="照明")
        result = ExtractionResult(
            title="品牌与单价测试",
            boxes=[box],
            circuits=[cir_in, cir_out],
            components=[Component(name="空开", spec="DZ47-63 C10/1P", quantity=1.0, used_in="AL1")],
            requirements=[],
            uncertainties=[],
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            path = tf.name

        try:
            # 指定目标品牌为 "正泰"
            build_workbook(result, "subtitle", path, layout="3_sheets", target_brand="正泰")
            wb = openpyxl.load_workbook(path)
            ws_det = wb["屏柜分项表"]

            # 查看分项表数据行 (Row 7 是箱头横幅, Row 8 是卡片表头, 数据从 row 9 开始)
            # Row 9: 进线隔离开关
            inc_name = ws_det.cell(row=9, column=2).value
            inc_price = float(ws_det.cell(row=9, column=6).value)
            inc_brand = ws_det.cell(row=9, column=8).value
            self.assertIn("隔离开关", inc_name)
            self.assertEqual(inc_brand, "正泰")
            # 验证下限价 86.50 已解禁：2P 隔离开关在正泰价格库中远低于 86.50
            self.assertLess(inc_price, 86.50)

            # Row 10: 分支出线断路器
            brk_price = float(ws_det.cell(row=10, column=6).value)
            brk_brand = ws_det.cell(row=10, column=8).value
            self.assertEqual(brk_brand, "正泰")
            # 验证下限价 15.08 已解禁：正泰 1P C10 采购价远低于 15.08
            self.assertLess(brk_price, 15.08)

            # 验证 DZ47 没有被写成电涌保护器
            all_dev_names = [ws_det.cell(row=r, column=2).value for r in range(9, 16) if ws_det.cell(row=r, column=2).value]
            self.assertNotIn("电涌保护器", all_dev_names)

            # 查找壳体外壳行
            shell_row = None
            for r in range(9, 16):
                if ws_det.cell(row=r, column=2).value == "壳体":
                    shell_row = r
                    break
            self.assertIsNotNone(shell_row)
            shell_price = float(ws_det.cell(row=shell_row, column=6).value)
            # 验证壳体下限价 457.06 已解禁：400*500*160 小型箱体测算价格低于 457.06
            self.assertLess(shell_price, 457.06)
        finally:
            if os.path.exists(path):
                os.remove(path)

    def test_cad_extractor_flexible_coordinates_and_empty_guard(self):
        """验证 CAD 提取移除 150000<=y<=300000 限制，任意坐标均可提取，且空图能生成明确存疑项。"""
        import ezdxf
        from extractor.cad_extractor import extract_cad_table_data

        # 1. 在 y = 500 (远小于 150000) 处放置配电箱和回路
        doc = ezdxf.new("R2010")
        msp = doc.modelspace()
        msp.add_text("01AP1 动力配电箱", dxfattribs={"insert": (100, 500), "height": 300})
        msp.add_text("WL1", dxfattribs={"insert": (100, 400), "height": 100})
        msp.add_text("C20/1P", dxfattribs={"insert": (200, 400), "height": 100})
        msp.add_text("BV-3x2.5", dxfattribs={"insert": (300, 400), "height": 100})
        msp.add_text("照明回路", dxfattribs={"insert": (400, 400), "height": 100})

        raw = extract_cad_table_data(doc)
        self.assertEqual(len(raw.boxes), 1)
        self.assertEqual(raw.boxes[0].code, "01AP1")
        self.assertEqual(len(raw.circuits), 1)
        self.assertEqual(raw.circuits[0].circuit_no, "WL1")

        # 2. 完全没有配电箱标头的图纸，验证记录明确存疑提示
        doc_empty = ezdxf.new("R2010")
        msp_empty = doc_empty.modelspace()
        msp_empty.add_text("无意义的说明文字", dxfattribs={"insert": (0, 0), "height": 100})

        raw_empty = extract_cad_table_data(doc_empty)
        self.assertEqual(len(raw_empty.boxes), 0)
        uncertainty_texts = [u.text for u in raw_empty.uncertainties]
        self.assertTrue(any("未检索到有效配电箱系统图标头" in t for t in uncertainty_texts))


if __name__ == "__main__":
    unittest.main()



