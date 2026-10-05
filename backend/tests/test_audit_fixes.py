import os
import unittest
import test_support  # noqa: F401
from fastapi.testclient import TestClient
from app import app, jobs, WORKDIR
import db
from extractor.schema import Box, Circuit, Component, ExtractionResult
from extractor.excel import build_workbook, _estimate_box_costs


class TestAuditFixes(unittest.TestCase):
    def setUp(self):
        db.init_db()
        self.client = TestClient(app)

    def test_invalid_token_cannot_spoof_tenant(self):
        """测试 1：防身份伪造——失效/伪造 Token 即使带上目标企业的 X-Tenant-ID，也无法越权访问。"""
        # 创建租户 victim 的私有任务
        victim_job_id = "test_victim_private_job"
        db.db_save_job({
            "job_id": victim_job_id,
            "tenant_id": "tenant_victim",
            "user_id": "user_victim",
            "status": "done",
            "project": "受害企业核心项目",
            "data": {"boxes": [{"code": "AP1"}], "circuits": [], "components": [], "requirements": []},
            "summary": {"title": "受害企业图纸", "boxes": 1, "circuits": 0, "components": 0}
        })

        try:
            # 攻击者使用伪造/无效的 token，并在请求头中试图指定 X-Tenant-ID 为 tenant_victim
            spoof_headers = {
                "Authorization": "Bearer tk_invalid_or_forged_token_12345",
                "X-Tenant-ID": "tenant_victim"
            }
            res = self.client.get(f"/api/jobs/{victim_job_id}", headers=spoof_headers)
            # 应当被 403 Forbidden 严格拦截，绝对无法跨租户越权窃取
            self.assertEqual(res.status_code, 403)
            self.assertIn("无权访问其他租户", res.json()["detail"])
        finally:
            conn = db._get_conn()
            with conn:
                conn.execute("DELETE FROM jobs WHERE id = ?", (victim_job_id,))

    def test_recalculate_components_after_restart_without_raw(self):
        """测试 2：一致性持久化恢复——重启后（无 raw 缓存），人工修改回路后 components 仍能联动重新组装。"""
        job_id = "test_job_restart_recalculate"
        # 模拟服务重启场景：jobs 字典中无 raw，仅有从数据库读出的 boxes 与 circuits
        db.db_save_job({
            "job_id": job_id,
            "tenant_id": "default",
            "status": "done",
            "filename": "测试图纸.pdf",
            "data": {
                "boxes": [{"code": "AL1", "name": "照明箱", "install": "挂墙", "size": "400*500"}],
                "circuits": [
                    {"box": "AL1", "circuit_no": "WL1", "breaker": "C16/1P", "cable": "BV-3x2.5", "load_name": "照明"}
                ],
                "components": [
                    {"name": "微型断路器", "spec": "C16/1P", "quantity": 1, "used_in": "AL1"}
                ],
                "requirements": [],
                "uncertainties": []
            }
        })

        try:
            # 清除内存 jobs 字典，模拟进程重启
            if job_id in jobs:
                del jobs[job_id]

            # 用户在前端修改回路：把断路器改成 C32/2P，并新增 WL2 回路
            update_payload = {
                "boxes": [{"code": "AL1", "name": "照明箱", "install": "挂墙", "size": "400*500"}],
                "circuits": [
                    {"box": "AL1", "circuit_no": "WL1", "breaker": "C32/2P", "cable": "BV-3x4", "load_name": "照明"},
                    {"box": "AL1", "circuit_no": "WL2", "breaker": "C20/1P", "cable": "BV-3x2.5", "load_name": "插座"}
                ],
                "components": [
                    # 传入旧的 components，验证服务端是否会自动推导重算
                    {"name": "微型断路器", "spec": "C16/1P", "quantity": 1, "used_in": "AL1"}
                ],
                "requirements": [],
                "uncertainties": [],
                "changes": [{"target": "WL1", "field": "breaker", "old": "C16/1P", "new": "C32/2P"}]
            }

            put_res = self.client.put(f"/api/jobs/{job_id}/data", json=update_payload)
            self.assertEqual(put_res.status_code, 200)
            saved_data = put_res.json()["data"]

            # 验证 components 已经联动重算：C16/1P 消失，出现了 C32/2P 和 C20/1P
            comp_specs = [c["spec"] for c in saved_data["components"]]
            self.assertNotIn("C16/1P", comp_specs)
            self.assertTrue(any("C32" in s for s in comp_specs))
            self.assertTrue(any("C20" in s for s in comp_specs))
        finally:
            conn = db._get_conn()
            with conn:
                conn.execute("DELETE FROM jobs WHERE id = ?", (job_id,))
            if job_id in jobs:
                del jobs[job_id]

    def test_box_quotation_includes_circuit_breakers(self):
        """测试 3：工业造价核算——正常回路断路器必须计入成套配电箱报价，金额不能仅算外壳。"""
        import tempfile
        import openpyxl

        box = Box(code="1AP1", name="动力配电箱", install="落地安装", location="车间A", size="800*1800*600")
        cir_inc = Circuit(box="1AP1", circuit_no="进线", breaker="NM1-125S/3300 100A", cable="", load_name="市电进线")
        cir_out1 = Circuit(box="1AP1", circuit_no="1WL1", breaker="C32/3P", cable="YJV-5x6", load_name="设备动力1")
        cir_out2 = Circuit(box="1AP1", circuit_no="1WL2", breaker="C20/1P", cable="BV-3x2.5", load_name="照明1")

        # 1. 验证成套估价函数内部正确提取断路器
        costs = _estimate_box_costs(box, [cir_inc, cir_out1, cir_out2], [])
        # 断路器元器件合计必须大于 0，且整体成套总单价必须大于 1800 元 (外壳+断路器+母排+工时)
        self.assertGreater(costs["comp_total"], 100.0)
        self.assertGreater(costs["unit_price"], 1800.0)

        # 2. 验证 Excel 导出分项明细中真实写入了开关器件行并计算了价格
        result = ExtractionResult(
            title="造价覆盖测试",
            boxes=[box],
            circuits=[cir_inc, cir_out1, cir_out2],
            components=[],
            requirements=[],
            uncertainties=[]
        )

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            path = tf.name

        try:
            build_workbook(result, "subtitle", path, layout="3_sheets")
            wb = openpyxl.load_workbook(path, data_only=False)
            ws_detail = wb["屏柜分项表"]
            # 查找断路器型号行
            found_breakers = []
            for r in range(7, 20):
                dev_name = ws_detail.cell(row=r, column=2).value
                dev_spec = ws_detail.cell(row=r, column=3).value
                unit_p = ws_detail.cell(row=r, column=6).value
                if dev_spec and any(k in str(dev_spec) for k in ["NM1", "C32", "C20"]):
                    found_breakers.append((dev_name, dev_spec, unit_p))
            self.assertEqual(len(found_breakers), 3)
            for _, _, p in found_breakers:
                self.assertGreater(float(p), 0.0)
        finally:
            if os.path.exists(path):
                os.remove(path)

    def test_prevent_cross_tenant_claim_exported_excel(self):
        """测试 4：跨企业数据防窃取——从磁盘恢复已导出 Excel 时，不能被其它企业随意认领。"""
        victim_job_id = "test_export_victim_99"
        # 1. 受害企业 tenant_A 导出并留有历史记录
        db.db_add_history({
            "job_id": victim_job_id,
            "filename": "绝密图纸.pdf",
            "project": "企业A专有项目",
            "title": "配电箱清单",
            "circuits": 5,
            "boxes": 1,
            "size": 1024
        }, tenant_id="tenant_A")

        # 2. 模拟磁盘上存在该任务的导出 Excel 文件
        xlsx_path = os.path.join(WORKDIR, f"{victim_job_id}.xlsx")
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "屏柜汇总表"
        ws.cell(row=1, column=1, value="【箱柜编号：AP1】")
        wb.save(xlsx_path)

        try:
            # 3. 租户 B 试图以自己的身份访问并认领该任务
            headers_tenant_b = {"X-Tenant-ID": "tenant_B"}
            res = self.client.get(f"/api/jobs/{victim_job_id}", headers=headers_tenant_b)
            # 恢复后原所属租户 tenant_A 保持生效，租户 B 的访问应被 403 拦截
            self.assertEqual(res.status_code, 403)
        finally:
            if os.path.exists(xlsx_path):
                os.remove(xlsx_path)
            conn = db._get_conn()
            with conn:
                conn.execute("DELETE FROM export_history WHERE job_id = ?", (victim_job_id,))
                conn.execute("DELETE FROM jobs WHERE id = ?", (victim_job_id,))
            if victim_job_id in jobs:
                del jobs[victim_job_id]

    def test_ai_usage_saved_and_restored_from_db(self):
        """测试 5：AI Token 消耗量与计费审计落盘——服务重启后 ai_usage 不丢失。"""
        job_id = "test_ai_usage_restore_job"
        test_usage = {
            "model": "deepseek-flash",
            "calls_count": 2,
            "prompt_tokens": 1200,
            "prompt_cache_hit_tokens": 800,
            "prompt_cache_miss_tokens": 400,
            "completion_tokens": 300,
            "total_tokens": 1500,
            "cost_in": 0.00056,
            "cost_out": 0.00060,
            "total_cost": 0.00116,
            "currency": "￥",
            "logs": []
        }
        db.db_save_job({
            "job_id": job_id,
            "tenant_id": "default",
            "status": "done",
            "filename": "测试图纸.pdf",
            "ai_usage": test_usage,
            "summary": {"title": "测试图纸", "boxes": 1, "circuits": 0, "components": 0}
        })

        try:
            loaded = db.db_get_job(job_id)
            self.assertIsNotNone(loaded)
            self.assertIn("ai_usage", loaded)
            self.assertEqual(loaded["ai_usage"].get("total_tokens"), 1500)
            self.assertEqual(loaded["ai_usage"].get("prompt_cache_hit_tokens"), 800)
            self.assertAlmostEqual(loaded["ai_usage"].get("total_cost"), 0.00116, places=4)
        finally:
            conn = db._get_conn()
            with conn:
                conn.execute("DELETE FROM jobs WHERE id = ?", (job_id,))

    def test_estimate_cost_deepseek_cache_hits(self):
        """测试 6：按价格表来源与生效日，对 DeepSeek 缓存分段计费。"""
        from extractor.vision import estimate_cost
        from datetime import datetime
        from zoneinfo import ZoneInfo
        # 2026-09-07 周一 08:00 上海时间，非高峰：命中 ¥0.02/M、未命中 ¥1/M、输出 ¥4/M。
        now = datetime(2026, 9, 7, 8, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
        res = estimate_cost("deepseek-flash", 1000, 500,
                            prompt_cache_hit_tokens=800, prompt_cache_miss_tokens=200,
                            now=now)
        self.assertAlmostEqual(res["cost_in"], 0.000216, places=8)
        self.assertAlmostEqual(res["cost_out"], 0.002, places=8)
        self.assertAlmostEqual(res["total_cost"], 0.002216, places=8)
        self.assertEqual(res["pricing_period"], "off_peak")
        self.assertTrue(res["pricing_available"])


if __name__ == "__main__":
    unittest.main()
