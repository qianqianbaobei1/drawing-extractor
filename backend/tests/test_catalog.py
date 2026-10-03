# -*- coding: utf-8 -*-
"""低压元器件物料库与平替推荐引擎单元测试。"""

import unittest

from extractor.catalog import (
    analyze_components_replacement,
    identify_brand,
    parse_component_spec,
    recommend_replacements,
)


class TestCatalogReplacement(unittest.TestCase):
    def test_identify_brand(self):
        self.assertEqual(identify_brand("施耐德 iC65N 2P C16A"), "施耐德")
        self.assertEqual(identify_brand("ABB S201-C16"), "ABB")
        self.assertEqual(identify_brand("西门子 5SY4116-7"), "西门子")
        self.assertEqual(identify_brand("正泰 NXB-63 C20/1P"), "正泰")
        self.assertEqual(identify_brand("德力西 CDB6i C32/3P"), "德力西")
        self.assertEqual(identify_brand("良信 NDB1-63 C10/1P"), "良信")
        self.assertEqual(identify_brand("MCB-63/C16A/1P"), "通用/国标")

    def test_parse_component_spec(self):
        res1 = parse_component_spec("微型断路器", "MCB-63/C16A/1P")
        self.assertEqual(res1["category"], "微型断路器 (MCB)")
        self.assertEqual(res1["poles"], "1P")
        self.assertEqual(res1["rated_amp"], 16.0)
        self.assertEqual(res1["curve"], "C")

        res2 = parse_component_spec("漏电开关", "RCB-63/C20A/1PN/30mA")
        self.assertIn("漏电", res2["category"])
        self.assertEqual(res2["poles"], "1P+N")
        self.assertEqual(res2["rated_amp"], 20.0)
        self.assertEqual(res2["leakage_ma"], 30)

        res3 = parse_component_spec("塑壳断路器", "MCCB-100/100A/3300")
        self.assertEqual(res3["category"], "塑壳断路器 (MCCB)")
        self.assertEqual(res3["poles"], "3P")
        self.assertEqual(res3["rated_amp"], 100.0)

    def test_recommend_replacements(self):
        # 施耐德微断替换为正泰
        rec_chint = recommend_replacements("微型断路器", "施耐德 iC65N 1P C16A", target_brand="正泰")
        self.assertEqual(rec_chint["original_brand"], "施耐德")
        self.assertEqual(rec_chint["target_brand"], "正泰")
        self.assertEqual(rec_chint["recommended_model"], "NXB-63 C16/1P")
        self.assertEqual(rec_chint["estimated_saving_pct"], 45)
        self.assertIn("16A", rec_chint["matching_notes"])

        # 施耐德漏电替换为德力西
        rec_delixi = recommend_replacements("漏电断路器", "施耐德 iDPN Vigi 1P+N C20A 30mA", target_brand="德力西")
        self.assertEqual(rec_delixi["target_brand"], "德力西")
        self.assertEqual(rec_delixi["recommended_model"], "CDB6LE-63 C20/1P+N 30mA")
        self.assertEqual(rec_delixi["estimated_saving_pct"], 45)

        # ABB 塑壳替换为良信
        rec_nader = recommend_replacements("塑壳断路器", "ABB XT1N 160 TMD 100A 3P", target_brand="良信")
        self.assertEqual(rec_nader["target_brand"], "良信")
        self.assertEqual(rec_nader["recommended_model"], "NDM1-125S/3300 100A 3P")
        self.assertEqual(rec_nader["estimated_saving_pct"], 35)

    def test_analyze_components_replacement(self):
        comps = [
            {"name": "微型断路器", "spec": "施耐德 iC65N C16A/1P", "quantity": 10, "unit": "只", "used_in": "WL1-WL10"},
            {"name": "塑壳断路器", "spec": "ABB XT1 100A 3P", "quantity": 1, "unit": "只", "used_in": "进线"},
            {"name": "配电箱体", "spec": "1AL1 600x400x160", "quantity": 1, "unit": "台", "used_in": "1AL1"},
        ]
        rep = analyze_components_replacement(comps, target_brand="正泰")
        self.assertEqual(rep["target_brand"], "正泰")
        self.assertEqual(rep["total_components"], 3)
        self.assertEqual(rep["total_quantity"], 12.0)
        self.assertTrue(rep["estimated_overall_saving_pct"] > 0)
        self.assertEqual(len(rep["items"]), 3)

    def test_excel_export_with_replacement(self):
        import tempfile
        import os
        import openpyxl
        from extractor.schema import ExtractionResult, Box, Circuit, Component
        from extractor.excel import build_workbook

        result = ExtractionResult(
            title="测试配电箱",
            boxes=[Box(code="1AL1", name="照明箱", ip_rating="IP30", install="明装", size="500x400x160", quantity=1)],
            circuits=[Circuit(circuit_no="WL1", breaker="MCB-63/C16A/1P", cable="BV-3x2.5", power_kw="1.5", load_name="照明")],
            components=[Component(name="微型断路器", spec="MCB-63/C16A/1P", unit="只", quantity=1, used_in="WL1")],
            requirements=[],
            uncertainties=[],
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = os.path.join(tmpdir, "test.xlsx")
            build_workbook(result, "副标题", out_file, target_brand="正泰", layout="all")
            self.assertTrue(os.path.exists(out_file))

            wb = openpyxl.load_workbook(out_file)
            self.assertIn("箱体清单", wb.sheetnames)
            self.assertIn("元器件汇总", wb.sheetnames)
            self.assertIn("国产化平替方案(正泰)", wb.sheetnames)

    def test_project_bom_excel_with_replacement(self):
        import tempfile
        import os
        import openpyxl
        from extractor.excel import build_project_bom_workbook

        mock_jobs = [
            {
                "box_code": "1AL1",
                "data": {
                    "boxes": [{"code": "1AL1", "name": "照明箱", "ip_rating": "IP30", "install": "明装", "size": "500x400x160", "quantity": 1}],
                    "circuits": [{"circuit_no": "WL1", "breaker": "MCB-63/C16A/1P"}],
                    "components": [{"name": "微型断路器", "spec": "MCB-63/C16A/1P", "unit": "只", "quantity": 2, "used_in": "1AL1"}],
                    "requirements": [{"item": "设计规范", "content": "按国标施工"}],
                }
            }
        ]

        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = os.path.join(tmpdir, "project_bom.xlsx")
            build_project_bom_workbook("测试项目", mock_jobs, out_file, target_brand="正泰")
            self.assertTrue(os.path.exists(out_file))

            wb = openpyxl.load_workbook(out_file)
            self.assertIn("全项目采购总清单(BOM)", wb.sheetnames)
            self.assertIn("集中采购平替(正泰)", wb.sheetnames)
            self.assertIn("配电箱成套设备台账", wb.sheetnames)
            self.assertIn("成套辅材与制造估算", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
